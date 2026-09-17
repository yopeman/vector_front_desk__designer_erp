"""Multi-format file generator for the Yope AI agent.

Generates files in /tmp/yope_generated/ and returns a unique filename.
Supported types: csv, json, txt, markdown, html, xlsx, pdf, docx
"""

import csv
import io
import json
import os
import uuid
from pathlib import Path

from app.config import settings

EXTENSION_MAP = {
    "csv": ".csv",
    "json": ".json",
    "txt": ".txt",
    "markdown": ".md",
    "md": ".md",
    "html": ".html",
    "xlsx": ".xlsx",
    "pdf": ".pdf",
    "docx": ".docx",
}

MIME_MAP = {
    "csv": "text/csv",
    "json": "application/json",
    "txt": "text/plain",
    "markdown": "text/markdown",
    "md": "text/markdown",
    "html": "text/html",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


def _ensure_dir() -> Path:
    d = Path(settings.generated_files_dir)
    d.mkdir(parents=True, exist_ok=True)
    return d


def _unique_name(ext: str) -> str:
    return f"{uuid.uuid4().hex}{ext}"


def generate_file(file_type: str, content: str, filename: str = "") -> dict:
    """Generate a file and return metadata.

    Args:
        file_type: one of csv, json, txt, markdown, md, html, xlsx, pdf, docx
        content: the raw text / JSON string / CSV text / HTML to write
        filename: optional human-friendly name (without extension)

    Returns:
        {"filename": str, "path": str, "mime_type": str, "size": int}
    """
    file_type = file_type.lower().strip()
    if file_type not in EXTENSION_MAP:
        raise ValueError(
            f"Unsupported file type '{file_type}'. "
            f"Supported: {', '.join(sorted(EXTENSION_MAP))}"
        )

    ext = EXTENSION_MAP[file_type]
    display_name = (filename or "report").strip()
    safe_name = "".join(c if c.isalnum() or c in "._-" else "_" for c in display_name)
    unique = _unique_name(ext)
    full_name = f"{safe_name}_{unique}"
    dir_path = _ensure_dir()
    file_path = dir_path / full_name

    if file_type == "csv":
        _write_csv(file_path, content)
    elif file_type == "json":
        _write_json(file_path, content)
    elif file_type in ("txt", "markdown", "md"):
        _write_text(file_path, content)
    elif file_type == "html":
        _write_html(file_path, content)
    elif file_type == "xlsx":
        _write_xlsx(file_path, content)
    elif file_type == "pdf":
        _write_pdf(file_path, content)
    elif file_type == "docx":
        _write_docx(file_path, content)

    size = file_path.stat().st_size
    return {
        "filename": full_name,
        "path": str(file_path),
        "mime_type": MIME_MAP[file_type],
        "size": size,
    }


def _write_csv(path: Path, content: str) -> None:
    """content is expected to be a JSON string of list-of-dicts or a raw CSV string."""
    try:
        data = json.loads(content)
        if isinstance(data, list) and data:
            with open(path, "w", newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=data[0].keys())
                writer.writeheader()
                writer.writerows(data)
            return
    except (json.JSONDecodeError, TypeError):
        pass
    # Fallback: write as raw text
    path.write_text(content, encoding="utf-8")


def _write_json(path: Path, content: str) -> None:
    try:
        data = json.loads(content)
        path.write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")
    except (json.JSONDecodeError, TypeError):
        path.write_text(content, encoding="utf-8")


def _write_text(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def _write_html(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def _write_xlsx(path: Path, content: str) -> None:
    """content is a JSON string of list-of-dicts or list-of-lists."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    try:
        data = json.loads(content)
    except (json.JSONDecodeError, TypeError):
        data = [[content]]

    if isinstance(data, list) and data:
        if isinstance(data[0], dict):
            headers = list(data[0].keys())
            ws.append(headers)
            for row in data:
                ws.append([row.get(h, "") for h in headers])
        elif isinstance(data[0], list):
            for row in data:
                ws.append(row)
        else:
            for item in data:
                ws.append([item])
    else:
        ws.append([str(data)])

    wb.save(str(path))
    wb.close()


def _write_pdf(path: Path, content: str) -> None:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)

    for line in content.split("\n"):
        pdf.cell(0, 7, txt=line, new_x="LMARGIN", new_y="NEXT")

    pdf.output(str(path))


def _write_docx(path: Path, content: str) -> None:
    from docx import Document

    doc = Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    doc.save(str(path))
